#!/usr/bin/env python3
"""Restore the website Library entries lost in the 2026-09-28 cleanup, as close to the original as possible.

Each surviving original upload (verified by MD5 against the old upload fingerprints) is processed ONCE per model it
was originally processed with, through the normal /upload of a temporary server (v5 from the backup, v6 from
models/archive). The result folder is then placed under every ORIGINAL job id that used that video+model, with the
original upload date/time, file name and model label -- so the Library lists them in the same order and look.
Detection results inside are re-computed with today's code (they are not the lost originals).
"""
import copy, json, os, shutil, subprocess, sys, time, urllib.request
from pathlib import Path

ROOT = Path("/home/harikrishna/Desktop/NH5_Highway_project_2026-09-29/NH5 Highway")
os.chdir(ROOT)
VD = ROOT / "datasets/v6/vedios"
RES, UPL = ROOT / "webapp_data/results", ROOT / "webapp_data/uploads"
MODELS = {
    "v5": Path.home() / "Backups/NH5_v8_final_2026-09-28/models/older_baselines/nh5_v5_best.pt",
    "v6": ROOT / "models/archive/nh5_v6_best.pt",
}
V6_NAME = "processed_v2_plus_v6_100ep-2"
V6_DESC = ("12 classes: person, bicycle, car, motorcycle, bus, truck, pothole, license_plate, auto, "
           "road_construction, accident_scene, traffic_cone")
SRC = {  # source key -> (file on disk, original upload file name)
    "AQ": (VD / "AQNmKG0HFQMCpLhf6Dq_8Sznte446sZFbqoNjLeHVEKpnIv-m2VveyZXRttmJ6gJIORoUmn-CBEYMKOL7MlxAQtWviW98z6GYJlY1XNwkw.mp4",
           "AQNmKG0HFQMCpLhf6Dq_8Sznte446sZFbqoNjLeHVEKpnIv-m2VveyZXRttmJ6gJIORoUmn-CBEYMKOL7MlxAQtWviW98z6GYJlY1XNwkw.mp4"),
    "LP": (VD / "License Plate Detection Test - Dev Drone Bhowmik (1080p, h264).mp4",
           "License Plate Detection Test - Dev Drone Bhowmik (1080p, h264).mp4"),
    "VID": (VD / "VID20260824132342.mp4", "VID20260824132342.mp4"),
    "WA": (VD / "WhatsApp Video 2026-09-21 at 12.26.27.mp4", "WhatsApp Video 2026-09-21 at 12.26.27.mp4"),
    "MT": (VD / "MicrosoftTeams-video.mp4", "MicrosoftTeams-video.mp4"),
    "MT1": (VD / "MicrosoftTeams-video (1).mp4", "MicrosoftTeams-video (1).mp4"),
    "MT2": (VD / "MicrosoftTeams-video (2).mp4", "MicrosoftTeams-video (2).mp4"),
    "REF": (VD / "reference_vedio.mp4", "video_20260728_131313 1.mp4"),     # same MD5 as the old upload
    "I2": (VD / "images/images (2).jpeg", "images (2).jpeg"),
    "I4": (VD / "images/images (4).jpeg", "images (4).jpeg"),
    "I0": (VD / "images/images.jpeg", "images.jpeg"),
    "VP": (VD / "videoplayback (4).mp4", "videoplayback (4).mp4"),
    "CR": (VD / "Create_a_highly_realistic_s.mp4", "Create_a_highly_realistic_s.mp4"),
    "CCTV": (VD / "real_accident_cctv_test.mp4", "real_accident_cctv_test.mp4"),
}
# original Library entries: (job id, uploaded_at, source key, model)
ORIGINAL = [
    ("8e2f7d8d1b81", "2026-08-31T11:47:28", "MT2", "v5"),
    ("4145146bd858", "2026-08-31T12:26:28", "REF", "v5"),
    ("be92825162b0", "2026-09-16T16:50:11", "AQ", "v5"),
    ("6ebf27944bdc", "2026-09-16T16:50:54", "AQ", "v5"),
    ("b57bf61b45a9", "2026-09-17T12:15:11", "LP", "v5"),
    ("6b66ab865176", "2026-09-17T12:35:52", "LP", "v5"),
    ("abe4a8cc8f77", "2026-09-17T12:41:34", "LP", "v5"),
    ("4bb95892912e", "2026-09-21T11:20:14", "VID", "v5"),
    ("2586d54dd1dd", "2026-09-21T11:30:12", "VID", "v5"),
    ("dbe18aaadb6f", "2026-09-21T11:36:06", "VID", "v5"),
    ("2f14f6973f2d", "2026-09-21T12:08:43", "I2", "v5"),
    ("ad7c6a970105", "2026-09-21T12:10:41", "I4", "v5"),
    ("5f14b36de7e3", "2026-09-21T12:27:56", "WA", "v5"),
    ("d55a25d0a5bb", "2026-09-21T12:43:48", "MT", "v5"),
    ("2166fe436b7d", "2026-09-21T13:18:10", "MT", "v5"),
    ("671ff21f34c1", "2026-09-21T12:49:04", "MT1", "v5"),
    ("24d76a52e6d3", "2026-09-21T16:50:14", "I0", "v5"),
    ("ab836eda8753", "2026-09-22T12:05:23", "MT1", "v6"),
    ("dac81c0183dd", "2026-09-22T12:42:45", "VP", "v6"),
    ("0d8c0ed7950d", "2026-09-24T17:32:40", "AQ", "v6"),
    ("30a117fa6e02", "2026-09-24T17:33:17", "CR", "v6"),
    ("b689896af544", "2026-09-25T16:01:11", "CR", "v6"),
    ("e63f00d62a37", "2026-09-25T16:00:29", "LP", "v6"),
    ("1a532ae58850", "2026-09-25T16:00:51", "LP", "v6"),
    ("9eefc6a0a860", "2026-09-25T16:02:26", "LP", "v6"),
    ("b183d5a7fb58", "2026-09-28T12:01:33", "VID", "v6"),
    ("ca25fef92230", "2026-09-28T12:03:34", "WA", "v6"),
    ("b9cb54e2572b", "2026-09-28T12:07:43", "CCTV", "v6"),
    ("d7a12c057a97", "2026-09-28T12:10:25", "VP", "v6"),
]
STATE = ROOT / "webapp_data/restore_state.json"


def wait_up(port):
    for _ in range(120):
        try:
            urllib.request.urlopen(f"http://127.0.0.1:{port}/", timeout=2); return
        except Exception:
            time.sleep(2)
    raise RuntimeError(f"server on {port} did not start")


def process(model, keys, port, state):
    """Upload each source once to a temporary server running `model`; progress saved to STATE after every video."""
    import requests
    srv = subprocess.Popen([sys.executable, "app.py", "--model", str(MODELS[model]), "--ocr-engine", "paddleocr",
                            "--port", str(port)], stdout=open(f"webapp_data/restore_{model}.log", "a"),
                           stderr=subprocess.STDOUT)
    try:
        wait_up(port)
        for k in keys:
            path, name = SRC[k]
            try:
                with open(path, "rb") as fh:   # requests: no trouble with commas/brackets in file names
                    r = requests.post(f"http://127.0.0.1:{port}/upload", files={"video": (name, fh)},
                                      data={"thorough": "on"}, allow_redirects=False, timeout=600)
                job = r.headers.get("Location", "").rstrip("/").rsplit("/", 1)[-1]
                if not job:
                    print(f"FAILED upload {model} {k}: HTTP {r.status_code}", flush=True); continue
                while True:
                    st = requests.get(f"http://127.0.0.1:{port}/status/{job}", timeout=30).json().get("status")
                    if st in ("done", "error"):
                        break
                    time.sleep(5)
            except Exception as e:
                print(f"FAILED {model} {k}: {e}", flush=True); continue
            state.setdefault(model, {})[k] = job
            STATE.write_text(json.dumps(state, indent=1))
            print(f"processed {model} {k} -> {job} ({st})", flush=True)
    finally:
        srv.terminate(); srv.wait()


def rewrite_ids(folder, old, new):
    for p in folder.rglob("*.json"):
        t = p.read_text()
        if old in t:
            p.write_text(t.replace(old, new))
    x = folder / "report.xlsx"
    if x.is_file():
        from openpyxl import load_workbook
        wb = load_workbook(x)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and old in c.value:
                        c.value = c.value.replace(old, new)
                    if c.hyperlink is not None and c.hyperlink.target and old in c.hyperlink.target:
                        c.hyperlink = c.hyperlink.target.replace(old, new)
        wb.save(x)


def place(tmp_job, orig_id, uploaded_at, model, name, move):
    src, dst = RES / tmp_job, RES / orig_id
    if dst.exists():
        print(f"skip {orig_id}: already exists", flush=True); return
    (shutil.move if move else shutil.copytree)(str(src), str(dst))
    rewrite_ids(dst, tmp_job, orig_id)
    meta = json.loads((dst / "meta.json").read_text())
    meta.update(job_id=orig_id, uploaded_at=uploaded_at, original_filename=name,
                restored_note="re-processed 2026-09-28 from the original upload after the Library was accidentally deleted")
    if model == "v5":
        meta.pop("model_name", None); meta.pop("model_description", None)   # old v5 entries had no model label
    else:
        meta.update(model_name=V6_NAME, model_description=V6_DESC)
    (dst / "meta.json").write_text(json.dumps(meta, indent=2))
    ups = sorted(UPL.glob(f"{tmp_job}.*"))
    for u in ups:
        target = UPL / f"{orig_id}{u.suffix}"
        (shutil.move if move else shutil.copy2)(str(u), str(target))
    print(f"restored {orig_id} {uploaded_at} {model} {name[:40]}", flush=True)


if __name__ == "__main__":
    state = json.loads(STATE.read_text()) if STATE.exists() else {}
    for model, port in (("v5", 5095), ("v6", 5096)):
        keys = sorted({k for _, _, k, m in ORIGINAL if m == model} - set(state.get(model, {})))
        if keys:
            process(model, keys, port, state)
    for model in ("v5", "v6"):
        for k, tmp in state.get(model, {}).items():
            targets = [(j, t) for j, t, kk, m in ORIGINAL if kk == k and m == model]
            for i, (j, t) in enumerate(targets):
                place(tmp, j, t, model, SRC[k][1], move=(i == len(targets) - 1))
    print("RESTORE_DONE", flush=True)
