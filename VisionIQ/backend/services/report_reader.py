"""Reads the Excel report NH5 writes for every processed video (webapp_data/results/<job>/report.xlsx).

Its "Vehicles" sheet uses VisionIQ's own column layout (ID, Timestamp (IST), Vehicle Type, Vehicle Number Plate,
Plate Confidence, ..., Plate Image Path) and the "Events" sheet lists accidents / fire / road work / behaviour
events. Rows keep their sheet row number so the ontology can point back to the exact spreadsheet row.
Parsed sheets are cached per file modification time."""
import threading
from pathlib import Path

from services.nh5_bridge import _local_url

_cache = {}
_lock = threading.Lock()

VEHICLE_COLUMNS = {
    'ID': 'id', 'Timestamp (IST)': 'timestamp', 'Vehicle Type': 'vehicleType', 'Vehicle Number Plate': 'plate',
    'Plate Confidence': 'plateConfidence', 'Speed (km/h)': 'speed', 'Speed Limit (km/h)': 'speedLimit',
    'Over Speed': 'overSpeed', 'Latitude': 'latitude', 'Longitude': 'longitude', 'Video Clip Path': 'videoClip',
    'Vehicle Image Path': 'vehicleImage', 'Plate Image Path': 'plateImage',
}
EVENT_COLUMNS = {
    'Event ID': 'id', 'Timestamp (IST)': 'timestamp', 'Video Time (s)': 'timeInVideoS', 'Event': 'label',
    'Severity': 'severity', 'Object': 'objectClass', 'Track ID': 'trackId', 'Vehicle Number Plate': 'plate',
    'Details': 'details', 'Snapshot': 'snapshot', 'Video Clip': 'clip',
}
URL_FIELDS = {'videoClip', 'vehicleImage', 'plateImage', 'snapshot', 'clip'}
NOT_A_PLATE = {'', 'not read', 'n/a', 'none', 'unknown'}


def _fraction(value):
    """'100%' / 100 / 0.97 -> 0..1 (None when not a number, e.g. 'n/a')."""
    if value is None or value == '':
        return None
    text = str(value).strip().rstrip('%')
    try:
        f = float(text)
    except ValueError:
        return None
    return f / 100.0 if (str(value).strip().endswith('%') or f > 1) else f


def _sheet_rows(ws, columns):
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    index = {columns[h]: i for i, h in enumerate(header) if h in columns}
    out = []
    for row_no, values in enumerate(rows, start=2):  # row 1 is the header
        if not values or all(v is None or v == '' for v in values):
            continue
        rec = {'row': row_no}
        for key, i in index.items():
            v = values[i] if i < len(values) else None
            rec[key] = _local_url(v) if key in URL_FIELDS and v else v
        if 'plateConfidence' in rec:
            rec['plateConfidence'] = _fraction(rec['plateConfidence'])
        plate = str(rec.get('plate') or '').strip()
        rec['plate'] = None if plate.lower() in NOT_A_PLATE else plate
        out.append(rec)
    return out


def read_report(job_dir):
    """{'file', 'vehicles': [...], 'events': [...]} from report.xlsx, or None when there is no report."""
    path = Path(job_dir) / 'report.xlsx'
    try:
        mtime = path.stat().st_mtime
    except OSError:
        return None
    with _lock:
        hit = _cache.get(str(path))
        if hit and hit[0] == mtime:
            return hit[1]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 -- unreadable / partially written report
        return None
    try:
        sheets = {ws.title: ws for ws in wb.worksheets}
        data = {
            'file': path.name,
            'vehicles': _sheet_rows(sheets['Vehicles'], VEHICLE_COLUMNS) if 'Vehicles' in sheets else [],
            'events': _sheet_rows(sheets['Events'], EVENT_COLUMNS) if 'Events' in sheets else [],
        }
    finally:
        wb.close()
    with _lock:
        _cache[str(path)] = (mtime, data)
    return data
