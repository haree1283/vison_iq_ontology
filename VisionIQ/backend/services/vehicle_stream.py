"""Live vehicle telemetry stream, stored-record queries, uploads and search (port of vehicleStream.js)."""
import math
import os
import random
import re
import threading
from datetime import datetime, timezone
from pathlib import Path

from config.database import BACKEND_DIR, db

STREAM_INTERVAL = int(os.getenv('STREAM_INTERVAL', '2500'))  # ms
DATA_DIRECTORY = Path(os.getenv('DATA_DIRECTORY', str(BACKEND_DIR / 'data'))).resolve()
NH44_FILE_NAME = 'NH44_vehicles_5000_merged_with_images.xlsx'

ALL_RECORDS_MODE = 'ALL_DATABASE_RECORDS'
MAX_STORED_VEHICLES = 20000  # /api/vehicles?limit= cap (the frontend loads every processed vehicle)
ALL_RECORDS_LABEL = 'All processed videos'

COLUMNS = ['csv_record_id', 'timestamp_ist', 'vehicle_type', 'vehicle_number_plate', 'plate_confidence',
           'vehicle_image', 'speed', 'speed_limit', 'over_speed', 'latitude', 'longitude', 'video_clip_path',
           'vehicle_image_path', 'plate_image_path', 'events', 'source_file', 'location', 'camera']

_state_lock = threading.RLock()
latest_stream_status = None
stream_paused = False
current_active_file_name = ALL_RECORDS_MODE
stream_offset = 0


def local_time_string():
    """Like JS new Date().toLocaleTimeString() (en-US): 4:05:09 PM"""
    return datetime.now().strftime('%I:%M:%S %p').lstrip('0')


def insert_sql(on_conflict='update'):
    placeholders = ', '.join(['%s'] * len(COLUMNS))
    sql = f"INSERT INTO vehicle_events ({', '.join(COLUMNS)}) VALUES ({placeholders})"
    if on_conflict == 'ignore':
        return sql.replace('INSERT', db.insert_ignore(), 1)
    return sql + db.upsert_suffix(
        ['source_file', 'csv_record_id'],
        ['speed', 'speed_limit', 'over_speed', 'timestamp_ist', 'vehicle_type', 'vehicle_number_plate',
         'plate_confidence', 'events', 'latitude', 'longitude'])


def _truthy(v):
    """JavaScript truthiness."""
    if v is None or v is False:
        return False
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v != 0 and not (isinstance(v, float) and math.isnan(v))
    if isinstance(v, str):
        return v != ''
    return True


def _first(*values):
    """JavaScript a || b || c."""
    for v in values:
        if _truthy(v):
            return v
    return values[-1] if values else None


def _js_str(v):
    if v is None:
        return 'null'
    if v is True:
        return 'true'
    if v is False:
        return 'false'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _js_number(row, key):
    """Number(row[key]) with JS semantics: missing -> NaN, null/'' -> 0."""
    if key not in row:
        return math.nan
    v = row[key]
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == '':
        return 0.0
    try:
        return float(s)
    except ValueError:
        return math.nan


def _finite(x):
    return isinstance(x, float) and math.isfinite(x)


def derive_location(lat, lon, id_num):
    if lat and lon:
        if lat < 17.25:
            return 'NH-44 - Shadnagar Interchange (KM 52)'
        if lat < 17.40:
            return 'NH-44 - Shamshabad Tollway (KM 18)'
        if lat < 17.50:
            return 'NH-44 - Hyderabad Ring Corridor'
        return 'NH-44 - Medchal North Gateway (KM 36)'
    locations = [
        'NH-44 - Shamshabad Tollway (KM 18)',
        'NH-44 - Shadnagar Interchange (KM 52)',
        'NH-44 - Hyderabad Ring Corridor',
        'NH-44 - Medchal North Gateway (KM 36)',
        'NH-44 - Jadcherla Express Point (KM 84)',
    ]
    return locations[(id_num or 0) % len(locations)]


def derive_camera(id_num):
    cams = [
        'CAM-NH44-01-SHAMSHABAD',
        'CAM-NH44-02-SHADNAGAR',
        'CAM-NH44-03-RINGROAD',
        'CAM-NH44-04-MEDCHAL',
        'CAM-NH44-05-JADCHERLA',
        'CAM-NH44-06-TOLLPLAZA',
    ]
    return cams[(id_num or 0) % len(cams)]


def canonical_vehicle_type(raw):
    if not raw:
        return 'Car'
    raw = str(raw)
    norm = raw.strip().lower()
    if 'truck' in norm or 'lorry' in norm or 'trailer' in norm:
        return 'Truck'
    if 'bus' in norm:
        return 'Bus'
    if 'auto' in norm or 'rickshaw' in norm:
        return 'Auto'
    if any(w in norm for w in ('bike', 'motorcycle', 'scooter', 'two wheeler')):
        return 'Bike'
    if 'van' in norm:
        return 'Van'
    if 'suv' in norm or 'jeep' in norm:
        return 'SUV'
    if 'tractor' in norm:
        return 'Tractor'
    if 'pedestrian' in norm:
        return 'Pedestrian'
    if 'train' in norm:
        return 'Train'
    if 'car' in norm or 'sedan' in norm or 'hatchback' in norm:
        return 'Car'
    return raw[:1].upper() + raw[1:].lower()


def _num_or(value, default):
    return float(value) if _truthy(value) else default


def to_client_record(row):
    digits = re.sub(r'\D', '', str(_first(row.get('id'), row.get('csv_record_id'), '')))
    id_num = int(digits) if digits else 0
    lat = float(row['latitude']) if _truthy(row.get('latitude')) else None
    lon = float(row['longitude']) if _truthy(row.get('longitude')) else None
    clean_type = canonical_vehicle_type(row.get('vehicle_type'))
    plate = row.get('vehicle_number_plate')
    # Node reported 0.95 for any missing confidence; a record whose plate was never read (NH5 detections)
    # reports 0 instead so the UI doesn't show a made-up 95% OCR confidence.
    default_conf = 0.95 if plate and plate != 'Not read' else 0
    confidence = _num_or(row.get('plate_confidence'), default_conf)

    return {
        'id': row.get('id'),
        'csvRecordId': row.get('csv_record_id'),
        'observationId': row.get('csv_record_id'),
        'timestampIst': row.get('timestamp_ist'),
        'timestamp': row.get('timestamp_ist'),
        'vehicleType': clean_type,
        'type': clean_type,
        'vehicleNumberPlate': plate,
        'numberPlate': plate,
        'plateConfidence': confidence,
        'confidence': confidence,
        'vehicleImage': row.get('vehicle_image'),
        'speed': _num_or(row.get('speed'), 0),
        'speedLimit': _num_or(row.get('speed_limit'), 60),
        'overSpeed': row.get('over_speed'),
        'isOverSpeed': str(row.get('over_speed')).lower() == 'yes',
        'latitude': lat,
        'longitude': lon,
        'location': row.get('location') or derive_location(lat, lon, id_num),
        'camera': row.get('camera') or derive_camera(id_num),
        'videoClipPath': row.get('video_clip_path'),
        'vehicleImagePath': row.get('vehicle_image_path') or row.get('vehicle_image'),
        'plateImagePath': row.get('plate_image_path'),
        'events': row.get('events'),
        'sourceFile': row.get('source_file'),
    }


def get_random_batch_size():
    """Random batch function: returns a random number of records between 1 and 20"""
    return random.randint(1, 20)


def _read_xlsx_rows(path):
    """First sheet -> list of dicts keyed by the header row (like XLSX.utils.sheet_to_json)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        keys = [str(h) if h is not None else f'__EMPTY_{i}' for i, h in enumerate(header)]
        out = []
        for values in rows:
            if values is None or all(v is None or v == '' for v in values):
                continue
            out.append({k: v for k, v in zip(keys, values) if v is not None and v != ''})
        return out
    finally:
        wb.close()


def seed_nh44_database_if_empty():
    """Ensure the 5,000 NH-44 records are stored in the database (when the xlsx is on disk)."""
    try:
        existing = db.query('SELECT COUNT(*) AS count FROM vehicle_events WHERE source_file = %s', (NH44_FILE_NAME,))
        if existing and existing[0]['count'] >= 5000:
            print(f"[Database] NH-44 dataset verified: {existing[0]['count']} records ready in the database.")
            return

        candidates = [BACKEND_DIR.parent / NH44_FILE_NAME, BACKEND_DIR / NH44_FILE_NAME, DATA_DIRECTORY / NH44_FILE_NAME]
        target = next((p for p in candidates if p.is_file()), None)
        if target is None:
            print(f'[Database Notice] {NH44_FILE_NAME} not found on disk, skipping auto-seed.')
            return

        print(f'[Database] Storing 5,000 records from {target} into the database...')
        rows = _read_xlsx_rows(target)
        values = []
        for idx, r in enumerate(rows, start=1):
            speed = float(r['Speed (km/h)']) if 'Speed (km/h)' in r else None
            speed_limit = float(r['Speed Limit (km/h)']) if 'Speed Limit (km/h)' in r else None
            over = str(r['Over Speed']) if r.get('Over Speed') else (
                'Yes' if speed and speed_limit and speed > speed_limit else 'No')
            values.append([
                f'NH44-{idx:05d}',
                str(r['Timestamp (IST)']) if r.get('Timestamp (IST)') else None,
                canonical_vehicle_type(r.get('Vehicle Type')),
                str(r['Vehicle Number Plate']) if r.get('Vehicle Number Plate') else None,
                float(r['Plate Confidence']) if 'Plate Confidence' in r else None,
                None,
                speed,
                speed_limit,
                over,
                float(r['Latitude']) if 'Latitude' in r else None,
                float(r['Longitude']) if 'Longitude' in r else None,
                str(r['Video Clip Path']) if r.get('Video Clip Path') else None,
                str(r['Vehicle Image Path']) if r.get('Vehicle Image Path') else None,
                None,
                str(r['Events']) if r.get('Events') else ('Overspeeding' if over == 'Yes' else 'Standard detection'),
                NH44_FILE_NAME,
                None,
                None,
            ])
        for i in range(0, len(values), 500):
            db.executemany(insert_sql('ignore'), values[i:i + 500])
        print(f'[Database] Successfully stored all {len(values)} records from {NH44_FILE_NAME}.')
    except Exception as error:  # noqa: BLE001
        print('[Database Seeding Error]:', error)


def is_all_mode(target):
    if not target:
        return True
    lower = str(target).lower().strip()
    return (lower in ('all', 'all_database_records', 'all database records')
            or 'all database' in lower or 'all records' in lower or 'all processed' in lower or '7,044' in lower)


def get_total_database_record_count():
    try:
        res = db.query('SELECT COUNT(*) AS count FROM vehicle_events')
        return int(res[0]['count'] or 0)
    except Exception:  # noqa: BLE001
        return 7044


def get_active_file_name():
    return ALL_RECORDS_LABEL if is_all_mode(current_active_file_name) else current_active_file_name


def fetch_random_vehicle_batch(forced_count=None, requested_file_name=None):
    """Fetch a random number of records (1 to 20) sequentially, looping through all records."""
    global stream_offset
    if isinstance(forced_count, (int, float)) and not isinstance(forced_count, bool) and forced_count > 0:
        count = int(min(forced_count, 50))
    else:
        count = get_random_batch_size()

    target_file = requested_file_name if requested_file_name is not None else current_active_file_name
    querying_all = is_all_mode(target_file)
    total_count = get_total_database_record_count()

    with _state_lock:
        if querying_all:
            if stream_offset >= total_count > 0:
                print(f'[VehicleStream] Stream reached end of database pool ({stream_offset}/{total_count}). '
                      'Looping back to offset 0.')
                stream_offset = 0
            rows = db.query('SELECT * FROM vehicle_events ORDER BY id ASC LIMIT %s OFFSET %s', (count, stream_offset))
            if not rows and total_count > 0:
                print(f'[VehicleStream] Loop wrap: returning to offset 0 across all {total_count} records.')
                stream_offset = 0
                rows = db.query('SELECT * FROM vehicle_events ORDER BY id ASC LIMIT %s OFFSET 0', (count,))
        else:
            like = f'%{target_file}%'
            res = db.query('SELECT COUNT(*) AS fileCount FROM vehicle_events WHERE source_file = %s OR source_file LIKE %s',
                           (target_file, like))
            file_total = int(res[0]['fileCount'] or 0) if res else 0
            if stream_offset >= file_total > 0:
                stream_offset = 0
            rows = db.query('SELECT * FROM vehicle_events WHERE source_file = %s OR source_file LIKE %s '
                            'ORDER BY id ASC LIMIT %s OFFSET %s', (target_file, like, count, stream_offset))
            if not rows and file_total > 0:
                stream_offset = 0
                rows = db.query('SELECT * FROM vehicle_events WHERE source_file = %s OR source_file LIKE %s '
                                'ORDER BY id ASC LIMIT %s OFFSET 0', (target_file, like, count))

        previous_offset = stream_offset
        stream_offset += len(rows)
        offset_now = stream_offset

    batch_time = datetime.now(timezone.utc).isoformat(timespec='milliseconds').replace('+00:00', 'Z')
    records = [{**to_client_record(r), '_streamSeq': previous_offset + i + 1, '_streamBatchTime': batch_time}
               for i, r in enumerate(rows)]
    return {'count': len(records), 'batchSize': count, 'records': records, 'offset': offset_now,
            'totalPool': total_count, 'isAllRecords': querying_all}


def store_uploaded_vehicle_records(file_name, rows, is_append=False):
    """Stores an uploaded vehicle dataset and includes it in the active database loop."""
    if not isinstance(rows, list) or not rows:
        raise ValueError('No records found in the uploaded file.')

    safe_file_name = str(file_name or 'uploaded_dataset.csv')[:255]
    base_name = re.sub(r'\.[^/.]+$', '', safe_file_name)
    now_str = datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')
    values = []
    for index, r in enumerate(rows, start=1):
        if not isinstance(r, dict):
            r = {}
        rec_id = _js_str(_first(r.get('id'), r.get('csvRecordId'), r.get('observationId'),
                                f'{base_name}-{index:05d}'))[:120]
        timestamp = _js_str(_first(r.get('timestampIst'), r.get('timestamp'), r.get('time'), now_str))[:80]
        vtype = canonical_vehicle_type(_js_str(_first(r.get('vehicleType'), r.get('type'), 'car')))
        plate = _js_str(_first(r.get('vehicleNumberPlate'), r.get('numberPlate'), r.get('plate'), ''))[:80]

        conf = 0.95
        if _truthy(r.get('plateConfidence')) or _truthy(r.get('confidence')):
            raw_conf = _js_number({'v': _first(r.get('plateConfidence'), r.get('confidence'))}, 'v')
            conf = min(max(raw_conf if _truthy(raw_conf) else 0.95, 0), 1)

        speed = _js_number(r, 'speed')
        speed = speed if _finite(speed) else float(random.randint(45, 84))
        speed_limit = _js_number(r, 'speedLimit')
        speed_limit = speed_limit if _finite(speed_limit) else 60.0
        is_over = _truthy(r.get('isOverSpeed')) or r.get('overSpeed') == 'Yes' or speed > speed_limit
        over_speed = 'Yes' if is_over else 'No'
        lat = _js_number(r, 'latitude')
        lat = lat if _finite(lat) else 17.385044
        lon = _js_number(r, 'longitude')
        lon = lon if _finite(lon) else 78.486671

        def path_or_none(key):
            v = r.get(key)
            if _truthy(v) and not str(v).startswith('data:'):
                return str(v)[:1000]
            return None

        events = _js_str(_first(r.get('events'), 'Overspeeding' if is_over else 'Standard detection'))[:255]
        values.append([rec_id, timestamp, vtype, plate, conf, None, speed, speed_limit, over_speed, lat, lon,
                       path_or_none('videoClipPath'), path_or_none('vehicleImagePath'),
                       path_or_none('plateImagePath'), events, safe_file_name, None, None])

    for i in range(0, len(values), 250):
        db.executemany(insert_sql('update'), values[i:i + 250])

    total_in_db = get_total_database_record_count()
    print(f'[Database] Successfully stored {len(rows)} records for uploaded file "{safe_file_name}". '
          f'Total records in database: {total_in_db}.')

    initial_batch = fetch_random_vehicle_batch(None, ALL_RECORDS_MODE)
    return {'success': True, 'totalStored': total_in_db, 'chunkStored': len(rows), 'fileName': ALL_RECORDS_LABEL,
            'initialBatch': initial_batch['records']}


def pause_stream():
    global stream_paused
    stream_paused = True
    return {'paused': True}


def resume_stream():
    global stream_paused
    stream_paused = False
    return {'paused': False}


def is_stream_paused():
    return stream_paused


def _broadcast_batch(socketio, batch, manual=False):
    global latest_stream_status
    socketio.emit('newVehicleBatch', batch['records'])
    status = {
        'type': 'batch',
        'fileName': get_active_file_name(),
        'batchCount': batch['count'],
        'totalPool': get_total_database_record_count(),
        'streamOffset': stream_offset,
        'timestamp': local_time_string(),
        'overspeeding': sum(1 for r in batch['records'] if r['isOverSpeed']),
        'isPaused': stream_paused,
    }
    if manual:
        status['manualTrigger'] = True
    latest_stream_status = status
    socketio.emit('vehicleStreamStatus', status)
    return status


def trigger_instant_batch(socketio):
    batch = fetch_random_vehicle_batch()
    if batch['records'] and socketio:
        _broadcast_batch(socketio, batch, manual=True)
    return batch


def run_vehicle_stream(socketio):
    seed_nh44_database_if_empty()
    live_total = get_total_database_record_count()
    print(f'[VehicleStream] Live telemetry streaming initialized across ALL {live_total} records in the database '
          f'(random batches 1-20 records every {STREAM_INTERVAL / 1000}s)')
    while True:
        try:
            if not stream_paused:
                batch = fetch_random_vehicle_batch()
                if batch['records'] and socketio:
                    status = _broadcast_batch(socketio, batch)
                    print(f"[VehicleStream] Broadcast random batch of {batch['count']} records "
                          f"(Offset: {stream_offset} / {status['totalPool']} database records)")
        except Exception as err:  # noqa: BLE001
            print('[VehicleStream Loop Error]:', err)
        socketio.sleep(STREAM_INTERVAL / 1000)


def get_stored_vehicles(limit=20, file_name=None):
    """Initial sample of stored vehicles (default 20)."""
    try:
        safe_limit = int(min(max(float(limit) if limit else 20, 1), MAX_STORED_VEHICLES))
    except (TypeError, ValueError):
        safe_limit = 20
    target_file = file_name if file_name is not None else current_active_file_name
    if is_all_mode(target_file):
        rows = db.query('SELECT * FROM vehicle_events ORDER BY id DESC LIMIT %s', (safe_limit,))
    else:
        rows = db.query('SELECT * FROM vehicle_events WHERE source_file = %s OR source_file LIKE %s '
                        'ORDER BY id ASC LIMIT %s', (target_file, f'%{target_file}%', safe_limit))
    return [to_client_record(r) for r in rows]


def get_vehicle_pool_stats(file_name=None):
    target_file = file_name if file_name is not None else current_active_file_name
    is_all = is_all_mode(target_file)
    totals_sql = ("SELECT COUNT(*) AS count, AVG(speed) AS avgSpeed, "
                  "SUM(CASE WHEN over_speed = 'Yes' THEN 1 ELSE 0 END) AS overspeedCount FROM vehicle_events")
    types_sql = 'SELECT vehicle_type, COUNT(*) AS count FROM vehicle_events'
    if is_all:
        total_res = db.query(totals_sql)
        types_res = db.query(types_sql + ' GROUP BY vehicle_type')
    else:
        where, params = ' WHERE source_file = %s OR source_file LIKE %s', (target_file, f'%{target_file}%')
        total_res = db.query(totals_sql + where, params)
        types_res = db.query(types_sql + where + ' GROUP BY vehicle_type', params)

    extra = db.query("SELECT COUNT(DISTINCT camera) AS cameras, COUNT(DISTINCT CASE WHEN vehicle_number_plate NOT IN "
                     "('', 'Not read') THEN vehicle_number_plate END) AS plates FROM vehicle_events"
                     + ('' if is_all else ' WHERE source_file = %s OR source_file LIKE %s'),
                     () if is_all else (target_file, f'%{target_file}%'))
    cameras = int(extra[0]['cameras'] or 0) if extra else 0
    plates = int(extra[0]['plates'] or 0) if extra else 0
    count = int(total_res[0]['count'] or 0) if total_res else 0
    avg_speed = round(float(total_res[0]['avgSpeed'] or 0)) if total_res else 0
    overspeed_total = int(total_res[0]['overspeedCount'] or 0) if total_res else 0
    return {
        'totalPool': count,
        'totalRecords': count,
        'averageSpeed': avg_speed,
        'avgSpeed': avg_speed,
        'overspeedTotal': overspeed_total,
        'overspeedRate': f'{round(overspeed_total / count * 100)}%' if count > 0 else '0%',
        'vehicleTypes': {r['vehicle_type']: r['count'] for r in types_res},
        'fileName': ALL_RECORDS_LABEL if is_all else target_file,
        'cameras': cameras,
        'plates': plates,
        'activeCorridor': 'NH5 Highway',
    }


def get_latest_stream_status():
    return latest_stream_status


def search_vehicles(search_term, limit=50):
    """Fast search by number plate, ID, type, or events; extracts plates / record IDs from natural language."""
    raw_term = str(search_term or '').strip()
    if not raw_term:
        return []

    plate_match = re.search(r'\b([a-zA-Z]{2}[-\s]?[0-9]{1,2}[-\s]?[a-zA-Z]{0,3}[-\s]?[0-9]{1,4})\b', raw_term, re.I)
    extracted_plate = re.sub(r'[^a-zA-Z0-9]', '', plate_match.group(1)) if plate_match else ''
    id_match = re.search(r'\b(NH44[-\s]?[0-9]+|OBS[-\s]?[0-9]+|ID[-\s]?[0-9]+)\b', raw_term, re.I)
    extracted_id = id_match.group(1) if id_match else ''
    clean_term = re.sub(r'[^a-zA-Z0-9]', '', raw_term)
    try:
        safe_limit = int(min(max(float(limit) if limit else 50, 1), 200))
    except (TypeError, ValueError):
        safe_limit = 50
    primary_plate = extracted_plate or clean_term

    try:
        rows = db.query(
            """SELECT * FROM vehicle_events
               WHERE REPLACE(REPLACE(vehicle_number_plate, ' ', ''), '-', '') LIKE %s
                  OR vehicle_number_plate LIKE %s
                  OR csv_record_id LIKE %s
                  OR vehicle_type LIKE %s
                  OR events LIKE %s
               ORDER BY
                 CASE
                   WHEN REPLACE(REPLACE(vehicle_number_plate, ' ', ''), '-', '') = %s THEN 0
                   WHEN REPLACE(REPLACE(vehicle_number_plate, ' ', ''), '-', '') LIKE %s THEN 1
                   ELSE 2
                 END,
                 id ASC LIMIT %s""",
            (f'%{primary_plate}%', f'%{raw_term}%', f'%{extracted_id or raw_term}%', f'%{raw_term}%',
             f'%{raw_term}%', primary_plate, f'%{primary_plate}%', safe_limit),
        )
        return [to_client_record(r) for r in rows]
    except Exception as error:  # noqa: BLE001
        print('[Database Search Error]:', error)
        return []
